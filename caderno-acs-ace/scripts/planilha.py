# -*- coding: utf-8 -*-
"""
planilha.py — Gera o XLSX do caderno.

Abas: Fontes e Metodologia · Série mensal e-Gestor · Composição por CBO · Resumo Setembro
      · Relação nominal (para cruzar com a lista de filiados/clientes).

Números vindos das fontes entram como VALOR; tudo o que é calculado entra como FÓRMULA
(ACS total, diferença, situação, matriz ano × mês) — assim o Juízo/perito vê a conta.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apuracao import MESES, LIMITE_DIVERGENCIA
from fontes import CBOS

FONTE = "Times New Roman"
CAB = PatternFill("solid", fgColor="D9D9D9")
FINO = Side(style="thin", color="808080")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def _cabecalho(ws, linha, titulos, larguras=None):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(name=FONTE, bold=True, size=10)
        c.fill = CAB
        c.border = BORDA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if larguras:
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _celula(ws, linha, col, valor, num=None, centro=False, wrap=False):
    c = ws.cell(row=linha, column=col, value=valor)
    c.font = Font(name=FONTE, size=10)
    c.border = BORDA
    c.alignment = Alignment(horizontal="center" if centro else "left", vertical="top", wrap_text=wrap)
    if num:
        c.number_format = num
    return c


def gerar_xlsx(dados, caminho):
    mun = dados["municipio"]
    wb = Workbook()

    # ------------------------------------------------------------------ Aba 1
    ws = wb.active
    ws.title = "Fontes e Metodologia"
    ws["A1"] = f"CADERNO DE PROVAS — ACS/ACE — {mun['nome'].upper()}/{mun['uf']}"
    ws["A1"].font = Font(name=FONTE, bold=True, size=12)
    info = [
        ("Município / UF", f"{mun['nome']} / {mun['uf']}"),
        ("Código IBGE (7 dígitos)", mun["ibge7"]),
        ("Código IBGE (6 dígitos, padrão CNES)", mun["ibge6"]),
        ("Código UF / Região", f"{mun['couf']} / {mun['coregiao']}"),
        ("Período", f"{dados['inicio']} a {dados['fim']} (AAAAMM)"),
        ("Data/hora da extração", dados["data_extracao"]),
        ("Regra de apuração", "ACS total = CBO 515105 + CBO 322255; conferência com qtAcsAtivoAb do e-Gestor; "
                              f"|diferença| > {LIMITE_DIVERGENCIA} = investigar"),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name=FONTE, bold=True, size=10)
        ws.cell(row=i, column=2, value=v).font = Font(name=FONTE, size=10)
    lin = 3 + len(info) + 1
    _cabecalho(ws, lin, ["Código", "Fonte", "Objeto", "Método", "URL exata", "Parâmetros", "Arquivo preservado",
                         "Bytes", "SHA-256", "Data/hora", "Situação"],
               [34, 16, 36, 8, 60, 40, 40, 10, 68, 30, 12])
    for r in dados["evidencias"]:
        lin += 1
        params = "; ".join(f"{k}={v}" for k, v in (r.get("parametros") or {}).items())
        vals = [r["codigo"], r["fonte"], r["descricao"], r["metodo"], r["url"], params, r["arquivo"],
                r["bytes"], r["sha256"], r["data_hora"], r["status"] if r["status"] == "ok" else f"FALHA: {r['erro']}"]
        for j, v in enumerate(vals, start=1):
            _celula(ws, lin, j, v, wrap=j in (3, 5, 6))
    ws.freeze_panes = ws.cell(row=3 + len(info) + 2, column=2)

    # ------------------------------------------------------------------ Aba 2
    ws2 = wb.create_sheet("Série mensal e-Gestor")
    _cabecalho(ws2, 1, ["Competência (AAAAMM)", "Mês/Ano", "ACS ativos (qtAcsAtivoAb)",
                        "ACS considerados p/ cobertura (qtAcsCoberturaAb)", "População (qtPopulacao)",
                        "Ano da população", "Cobertura ACS % (pcCoberturaAcsAb)"],
               [16, 11, 16, 20, 16, 12, 16])
    serie = dados["egestor"]
    for i, l in enumerate(serie, start=2):
        _celula(ws2, i, 1, l["nuComp"], centro=True)            # texto: preserva o formato AAAAMM
        _celula(ws2, i, 2, f"{MESES[int(l['nuComp'][4:]) - 1]}/{l['nuComp'][:4]}", centro=True)
        _celula(ws2, i, 3, l["qtAcsAtivoAb"], "0", centro=True)
        _celula(ws2, i, 4, l["qtAcsCoberturaAb"], "0", centro=True)
        _celula(ws2, i, 5, l["qtPopulacao"], "#,##0", centro=True)
        _celula(ws2, i, 6, l["anoPopulacao"], centro=True)
        _celula(ws2, i, 7, l["pcCoberturaAcsAb"], "0.00", centro=True)
    ult = max(2, len(serie) + 1)
    # Matriz ano × mês calculada por FÓRMULA a partir da lista acima
    col0 = 9
    ws2.cell(row=1, column=col0, value="Matriz ano × mês — ACS ativos (fórmula sobre as colunas A e C)").font = \
        Font(name=FONTE, bold=True, size=10)
    titulos = ["Ano"] + MESES
    for j, t in enumerate(titulos):
        c = ws2.cell(row=2, column=col0 + j, value=t)
        c.font = Font(name=FONTE, bold=True, size=10)
        c.fill = CAB
        c.border = BORDA
        c.alignment = Alignment(horizontal="center")
        ws2.column_dimensions[get_column_letter(col0 + j)].width = 8
    for i, ano in enumerate(range(int(dados["inicio"][:4]), int(dados["fim"][:4]) + 1), start=3):
        _celula(ws2, i, col0, ano, centro=True)
        for m in range(1, 13):
            comp = f"{ano}{m:02d}"
            f = f'=IFERROR(INDEX($C$2:$C${ult},MATCH("{comp}",$A$2:$A${ult},0)),"")'
            _celula(ws2, i, col0 + m, f, "0", centro=True)
    ws2.freeze_panes = "A2"

    # ------------------------------------------------------------------ Aba 3
    ws3 = wb.create_sheet("Composição por CBO")
    _cabecalho(ws3, 1, ["Competência (AAAAMM)", "Mês/Ano", "Papel", "CBO 515105 ACS", "CBO 322255 Técnico em ACS",
                        "ACS total (fórmula D+E)", "CBO 515140 ACE", "e-Gestor ACS ativos",
                        "Diferença (F−H)", "Situação", "Evidências CNES (códigos)"],
               [14, 10, 22, 12, 14, 14, 12, 12, 12, 26, 70])
    linha_por_comp = {}
    for i, q in enumerate(dados["quadro"], start=2):
        linha_por_comp[q["comp"]] = i
        _celula(ws3, i, 1, q["comp"], centro=True)
        _celula(ws3, i, 2, q["rotulo"], centro=True)
        _celula(ws3, i, 3, q["papel"])
        _celula(ws3, i, 4, q["cbo515105"], "0", centro=True)
        _celula(ws3, i, 5, q["cbo322255"], "0", centro=True)
        _celula(ws3, i, 6, f'=IF(COUNT(D{i}:E{i})<2,"",D{i}+E{i})', "0", centro=True)
        _celula(ws3, i, 7, q["ace515140"], "0", centro=True)
        _celula(ws3, i, 8, q["egestor_ativo"], "0", centro=True)
        _celula(ws3, i, 9, f'=IF(OR(F{i}="",H{i}=""),"",F{i}-H{i})', "0", centro=True)
        _celula(ws3, i, 10, f'=IF(F{i}="","Não apurado (CNES)",IF(H{i}="","Sem dado no e-Gestor",'
                            f'IF(ABS(I{i})>{LIMITE_DIVERGENCIA},"Divergência > 2 — investigar","Confere")))')
        _celula(ws3, i, 11, " · ".join(e for e in q["evidencias"] if e), wrap=True)
    ws3.freeze_panes = "C2"

    # ------------------------------------------------------------------ Aba 4
    ws4 = wb.create_sheet("Resumo Setembro")
    _cabecalho(ws4, 1, ["Exercício", "Competência", "ACS (515105)", "Técnico em ACS (322255)", "ACS total",
                        "ACE (515140)", "e-Gestor ACS ativos", "Diferença", "Situação"],
               [10, 12, 12, 14, 10, 10, 12, 10, 26])
    lin = 1
    for q in dados["quadro"]:
        if not q["comp"].endswith("09"):
            continue
        lin += 1
        o = linha_por_comp[q["comp"]]
        ref = "'Composição por CBO'!"
        _celula(ws4, lin, 1, int(q["comp"][:4]), centro=True)
        _celula(ws4, lin, 2, q["rotulo"], centro=True)
        for col, src in zip(range(3, 10), "DEFGHIJ"):
            # Célula vazia na origem continua vazia aqui (sem virar "0")
            _celula(ws4, lin, col, f'=IF({ref}{src}{o}="","",{ref}{src}{o})', "0", centro=True)
    if lin == 1:
        ws4["A2"] = "Nenhuma competência de setembro apurada no período."

    # ------------------------------------------------------------------ Aba 5
    ws5 = wb.create_sheet("Relação nominal")
    _cabecalho(ws5, 1, ["CBO", "Ocupação", "Nº", "Nome (posição da data da consulta)"], [10, 36, 6, 50])
    lin = 1
    for cbo, nom in dados["nominal"].items():
        if not nom.get("nomes"):
            continue
        for k, n in enumerate(nom["nomes"], start=1):
            lin += 1
            _celula(ws5, lin, 1, cbo, centro=True)
            _celula(ws5, lin, 2, CBOS[cbo])
            _celula(ws5, lin, 3, k, centro=True)
            _celula(ws5, lin, 4, n)
    ws5.freeze_panes = "A2"

    wb.save(caminho)
    return caminho
