# -*- coding: utf-8 -*-
"""
conferencia.py — Conferência obrigatória antes de entregar.

  1. Converte o DOCX em PDF (LibreOffice; se não houver, tenta o Word via docx2pdf).
  2. Renderiza cada página do PDF em imagem e monta folhas de contato (6 páginas por imagem)
     para inspeção visual — página em branco, tabela espremida, caixa partida, rótulo cortado.
  3. Detecta automaticamente páginas em branco.
  4. Confere se os marcadores (bookmarks) do PDF existem e apontam para as 8 seções.
  5. Recalcula o XLSX no LibreOffice e procura erros de fórmula (#REF!, #VALUE!, #NAME? ...),
     comparando os totais calculados pelas fórmulas com os números apurados.

Nada disso substitui OLHAR as folhas de contato — o relatório só aponta onde olhar primeiro.
"""

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

SECOES_ESPERADAS = ["1 ", "2 ", "3 ", "4 ", "5 ", "6 ", "7 ", "8 "]
ERROS_FORMULA = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")


def achar_soffice():
    for nome in ("soffice", "libreoffice"):
        c = shutil.which(nome)
        if c:
            return c
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(c):
            return c
    return None


def _soffice_converter(arquivo, destino, filtro):
    soffice = achar_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice não encontrado")
    perfil = Path(tempfile.mkdtemp(prefix="lo_perfil_"))   # perfil isolado: evita travar se o LO estiver aberto
    cmd = [soffice, f"-env:UserInstallation={perfil.as_uri()}", "--headless", "--norestore",
           "--convert-to", filtro, "--outdir", str(destino), str(arquivo)]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=240)
    finally:
        shutil.rmtree(perfil, ignore_errors=True)


def docx_para_pdf(docx, pasta):
    """Devolve o caminho do PDF ou levanta exceção com a explicação."""
    docx, pasta = Path(docx), Path(pasta)
    pdf = pasta / (docx.stem + ".pdf")
    try:
        _soffice_converter(docx, pasta, 'pdf:writer_pdf_Export:{"ExportBookmarks":{"type":"boolean","value":"true"}}')
        if pdf.exists():
            return pdf
    except Exception as e:
        erro_lo = e
    else:
        erro_lo = "LibreOffice não gerou o PDF"
    try:                                   # Plano B no Windows: Microsoft Word
        from docx2pdf import convert
        convert(str(docx), str(pdf))
        if pdf.exists():
            return pdf
    except Exception as e2:
        raise RuntimeError(f"Sem conversor DOCX→PDF (LibreOffice: {erro_lo}; Word/docx2pdf: {e2}). "
                           "Instale o LibreOffice ou rode: pip install docx2pdf") from e2
    raise RuntimeError("Conversão DOCX→PDF falhou sem mensagem.")


def renderizar_paginas(pdf, pasta, dpi=70):
    """PDF → PNG por página. Usa pypdfium2 (pip) e, na falta, pdftoppm (poppler)."""
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    imagens = []
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(pdf))
        for i in range(len(doc)):
            img = doc[i].render(scale=dpi / 72).to_pil()
            destino = pasta / f"pagina_{i + 1:03d}.png"
            img.save(destino)
            imagens.append(destino)
        doc.close()
        return imagens
    except ImportError:
        pass
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", str(dpi), str(pdf), str(pasta / "pagina")], check=True, timeout=240)
        return sorted(pasta.glob("pagina*.png"))
    raise RuntimeError("Sem renderizador de PDF (instale: pip install pypdfium2).")


def pagina_em_branco(img_path, limiar=0.004):
    """Página com menos de 0,4% de pixels escuros = em branco (só o número da página, por exemplo)."""
    from PIL import Image
    im = Image.open(img_path).convert("L")
    # Ignora a faixa do rodapé (número da página)
    w, h = im.size
    corpo = im.crop((0, 0, w, int(h * 0.90)))
    hist = corpo.histogram()
    escuros = sum(hist[:200])
    return escuros / (corpo.size[0] * corpo.size[1]) < limiar


def folhas_de_contato(imagens, pasta, por_folha=6, largura_mini=420):
    """Junta as páginas em grades 3×2 com o número de cada página — para olhar rápido."""
    from PIL import Image, ImageDraw
    pasta = Path(pasta)
    folhas = []
    for k in range(0, len(imagens), por_folha):
        grupo = imagens[k:k + por_folha]
        minis = []
        for p in grupo:
            im = Image.open(p).convert("RGB")
            esc = largura_mini / max(im.size)
            minis.append(im.resize((int(im.size[0] * esc), int(im.size[1] * esc))))
        cel_w, cel_h = largura_mini, largura_mini
        folha = Image.new("RGB", (cel_w * 3 + 40, cel_h * 2 + 60), "#9a9a9a")
        d = ImageDraw.Draw(folha)
        for i, (mini, p) in enumerate(zip(minis, grupo)):
            x = 10 + (i % 3) * (cel_w + 10)
            y = 30 + (i // 3) * (cel_h + 20)
            folha.paste(mini, (x + (cel_w - mini.size[0]) // 2, y))
            d.text((x + 4, y - 16), f"pág. {k + i + 1}", fill="white")
        destino = pasta / f"folha_contato_{k // por_folha + 1:02d}.png"
        folha.save(destino)
        folhas.append(destino)
    return folhas


def conferir_marcadores(pdf):
    """Lista os marcadores do PDF e verifica as 8 seções obrigatórias."""
    from pypdf import PdfReader
    leitor = PdfReader(str(pdf))
    titulos = []

    def andar(itens, nivel=0):
        for it in itens:
            if isinstance(it, list):
                andar(it, nivel + 1)
            else:
                try:
                    pag = leitor.get_destination_page_number(it) + 1
                except Exception:
                    pag = None
                titulos.append((nivel, it.title, pag))
    try:
        andar(leitor.outline)
    except Exception:
        pass
    primeiro_nivel = [t for t in titulos if t[0] == 0]
    faltando = [s.strip() for s in SECOES_ESPERADAS
                if not any(t[1].strip().startswith(s) or t[1].strip().startswith(s.strip() + " ") for t in primeiro_nivel)]
    sem_pagina = [t[1] for t in titulos if t[2] is None]
    return {"paginas": len(leitor.pages), "marcadores": titulos, "secoes_faltando": faltando,
            "marcadores_sem_destino": sem_pagina}


def recalcular_xlsx(xlsx, quadro):
    """
    Recalcula no LibreOffice (cópia temporária) e procura erros de fórmula.
    Também confere: ACS total da fórmula == 515105 + 322255 apurado em Python.
    """
    from openpyxl import load_workbook
    tmp = Path(tempfile.mkdtemp(prefix="recalc_"))
    try:
        _soffice_converter(xlsx, tmp, "xlsx")
        recalc = tmp / Path(xlsx).name
        wb = load_workbook(recalc, data_only=True)
    except Exception as e:
        shutil.rmtree(tmp, ignore_errors=True)
        return {"ok": None, "erro": f"não foi possível recalcular ({e}) — abrir no Excel e conferir manualmente"}
    erros = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERROS_FORMULA):
                    erros.append(f"{ws.title}!{c.coordinate} = {c.value}")
    divergentes = []
    ws = wb["Composição por CBO"]
    esperado = {q["comp"]: q for q in quadro}
    for row in ws.iter_rows(min_row=2):
        comp = row[0].value
        if comp in esperado:
            q = esperado[comp]
            calc_total = row[5].value
            calc_dif = row[8].value
            exp_total = q["acs_total"] if q["acs_total"] is not None else ""
            exp_dif = q["diferenca"] if q["diferenca"] is not None else ""
            def norm(v):
                return "" if v in (None, "") else int(round(float(v)))
            if norm(calc_total) != norm(exp_total) or norm(calc_dif) != norm(exp_dif):
                divergentes.append(f"{comp}: planilha total={calc_total}, dif={calc_dif} × apurado total={exp_total}, dif={exp_dif}")
    shutil.rmtree(tmp, ignore_errors=True)
    return {"ok": not erros and not divergentes, "erros_formula": erros, "divergencias_calculo": divergentes}


def conferir_tudo(docx, xlsx, quadro, pasta_saida):
    """Executa toda a conferência e devolve um relatório (dict)."""
    pasta_saida = Path(pasta_saida)
    # Limpa a conferência anterior (no --reprocessar, páginas antigas confundiriam a inspeção)
    shutil.rmtree(pasta_saida / "paginas", ignore_errors=True)
    for antiga in pasta_saida.glob("folha_contato_*.png"):
        antiga.unlink(missing_ok=True)
    rel = {"pdf": None, "paginas_em_branco": [], "folhas_contato": [], "marcadores": None, "xlsx": None, "erros": []}
    try:
        pdf = docx_para_pdf(docx, Path(docx).parent)
        rel["pdf"] = str(pdf)
        pasta_img = pasta_saida / "paginas"
        imgs = renderizar_paginas(pdf, pasta_img)
        rel["paginas_em_branco"] = [i + 1 for i, p in enumerate(imgs) if pagina_em_branco(p)]
        rel["folhas_contato"] = [str(f) for f in folhas_de_contato(imgs, pasta_saida)]
        rel["marcadores"] = conferir_marcadores(pdf)
    except Exception as e:
        rel["erros"].append(f"PDF/renderização: {e}")
    try:
        rel["xlsx"] = recalcular_xlsx(xlsx, quadro)
    except Exception as e:
        rel["erros"].append(f"XLSX: {e}")
    return rel
