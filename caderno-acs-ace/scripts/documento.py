# -*- coding: utf-8 -*-
"""
documento.py — Gera o Caderno de Provas em DOCX no padrão ABNT do escritório.

Times New Roman 12 · entrelinha 1,5 · A4 · margens sup. 3 / esq. 3 / inf. 2 / dir. 2 cm ·
justificado · rodapé paginado · títulos como "Título 1/2" (viram marcadores no PDF).

Estrutura (8 seções):
 1 Objeto e identificação do ente
 2 Metodologia, fontes e cadeia de obtenção
 3 Quadro-síntese — competência de setembro de cada exercício
 4 Série histórica mensal (paisagem)
 5 Composição por CBO + gráfico + "Achado relevante"
 6 Relação nominal por CBO (duas colunas, numerada)
 7 Roteiro de conferência
 8 Nota metodológica, limites e relação de fontes

Tom: impessoal, sem adjetivação, sem conclusão jurídica — o caderno é prova, não petição.
"""

import math
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from apuracao import MESES, LIMITE_DIVERGENCIA, fmt_dif, rotulo_comp
from fontes import CBOS, URL_CNES_CBO, URL_CNES_NOMINAL, URL_EGESTOR_PAINEL

FONTE = "Times New Roman"
MESES_EXTENSO = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
                 "setembro", "outubro", "novembro", "dezembro"]


# =============================================================================
# Estilos e utilidades de baixo nível
# =============================================================================
def _fixar_fonte(estilo_ou_run, nome=FONTE):
    """Força a fonte em todos os 'scripts' e remove a fonte de tema (senão o Word ignora)."""
    rpr = estilo_ou_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(att), nome)


def _configurar_estilos(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.size = Pt(12)
    _fixar_fonte(normal)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for nome, caps in (("Heading 1", True), ("Heading 2", False)):
        h = st[nome]
        _fixar_fonte(h)
        h.font.size = Pt(12)
        h.font.bold = True
        h.font.italic = False
        h.font.all_caps = caps
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing = 1.5
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Estilo para texto de tabela: simples, 10 pt, sem recuo
    tab = st.add_style("Tabela", 1)
    tab.base_style = normal
    tab.font.size = Pt(10)
    tab.paragraph_format.line_spacing = 1.0
    tab.paragraph_format.space_after = Pt(0)
    tab.paragraph_format.first_line_indent = Cm(0)
    tab.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Legenda (título acima / fonte abaixo de tabelas e gráficos — ABNT)
    leg = st.add_style("Legenda ABNT", 1)
    leg.base_style = normal
    leg.font.size = Pt(10)
    leg.paragraph_format.line_spacing = 1.0
    leg.paragraph_format.space_before = Pt(6)
    leg.paragraph_format.space_after = Pt(3)
    leg.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "List Bullet" in [s.name for s in st]:
        lb = st["List Bullet"]
        _fixar_fonte(lb)
        lb.font.size = Pt(12)


def _margens(secao, paisagem=False):
    if paisagem:
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = Cm(29.7), Cm(21.0)
    else:
        secao.orientation = WD_ORIENT.PORTRAIT
        secao.page_width, secao.page_height = Cm(21.0), Cm(29.7)
    secao.top_margin = Cm(3)
    secao.left_margin = Cm(3)
    secao.bottom_margin = Cm(2)
    secao.right_margin = Cm(2)
    secao.header_distance = Cm(1.25)
    secao.footer_distance = Cm(1.0)


def _campo_pagina(paragrafo):
    """Insere o campo PAGE (número da página) no rodapé."""
    run = paragrafo.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)
    run.font.size = Pt(10)


def _sombrear(celula, cor):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def _linha_nao_quebra(linha, cabecalho=False):
    trpr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trpr.append(el)
    if cabecalho:
        h = OxmlElement("w:tblHeader")
        trpr.append(h)


def _largura_fixa(tabela):
    tblpr = tabela._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)


class Caderno:
    """Pequena camada sobre o python-docx com os blocos do caderno."""

    def __init__(self):
        self.doc = Document()
        _configurar_estilos(self.doc)
        s = self.doc.sections[0]
        _margens(s)
        rod = s.footer.paragraphs[0]
        rod.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _campo_pagina(rod)
        self.n_tabela = 0
        self.n_grafico = 0
        self.paisagem = False

    # ---------------- blocos de texto
    def titulo(self, texto, nivel=1):
        return self.doc.add_heading(texto, level=nivel)

    def p(self, texto, recuo=True, negrito_ate=None, alinhamento=None, tamanho=None):
        """Parágrafo. negrito_ate='Fonte A —' deixa em negrito o início até esse trecho."""
        par = self.doc.add_paragraph()
        if recuo:
            par.paragraph_format.first_line_indent = Cm(1.25)
        if alinhamento is not None:
            par.alignment = alinhamento
        if negrito_ate and texto.startswith(negrito_ate):
            r = par.add_run(negrito_ate)
            r.bold = True
            texto = texto[len(negrito_ate):]
        r = par.add_run(texto)
        if tamanho:
            for run in par.runs:
                run.font.size = Pt(tamanho)
        return par

    def marcadores(self, itens):
        for it in itens:
            par = self.doc.add_paragraph(style="List Bullet")
            par.add_run(it)

    def passos(self, itens):
        """Lista numerada manual (numeração estável em Word e LibreOffice)."""
        for i, it in enumerate(itens, start=1):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            par.add_run(f"{i}. ").bold = True
            par.add_run(it)

    def legenda(self, texto, acima=True):
        par = self.doc.add_paragraph(style="Legenda ABNT")
        if acima:
            par.paragraph_format.keep_with_next = True
            par.add_run(texto).bold = True
        else:
            par.add_run(texto)
        return par

    # ---------------- tabelas
    def tabela(self, cabecalho, linhas, larguras_cm, tamanho=10, alinhar=None, titulo=None, fonte=None,
               mono_cols=()):
        """
        alinhar: string com 'E' (esquerda) / 'C' (centro) por coluna.
        mono_cols: colunas em Courier New (hashes).
        Linhas não quebram entre páginas; cabeçalho repete em cada página.
        """
        if titulo:
            self.n_tabela += 1
            self.legenda(f"Tabela {self.n_tabela} — {titulo}")
        n = len(cabecalho)
        alinhar = alinhar or ("E" * n)
        t = self.doc.add_table(rows=1, cols=n)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _largura_fixa(t)
        for j, w in enumerate(larguras_cm):
            t.columns[j].width = Cm(w)          # tblGrid — o LibreOffice lê daqui
        cab = t.rows[0]
        _linha_nao_quebra(cab, cabecalho=True)
        for j, txt in enumerate(cabecalho):
            c = cab.cells[j]
            c.width = Cm(larguras_cm[j])
            _sombrear(c, "D9D9D9")
            par = c.paragraphs[0]
            par.style = self.doc.styles["Tabela"]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(str(txt))
            r.bold = True
            r.font.size = Pt(tamanho)
        for linha in linhas:
            row = t.add_row()
            _linha_nao_quebra(row)
            for j, val in enumerate(linha):
                c = row.cells[j]
                c.width = Cm(larguras_cm[j])
                par = c.paragraphs[0]
                par.style = self.doc.styles["Tabela"]
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER if alinhar[j] == "C" else WD_ALIGN_PARAGRAPH.LEFT
                partes = val if isinstance(val, (list, tuple)) else [val]
                for k, parte in enumerate(partes):
                    r = par.add_run("—" if parte is None else str(parte))
                    r.font.size = Pt(tamanho)
                    if j in mono_cols:
                        r.font.name = "Courier New"
                        _fixar_fonte(r, "Courier New")
                    if k < len(partes) - 1:
                        r.add_break(WD_BREAK.LINE)
        # Tabela curta: fica inteira na página e presa à linha "Fonte".
        # Tabela longa: só a última linha prende a "Fonte" (evita "Fonte" sozinha na página seguinte).
        presas = t.rows[1:] if len(t.rows) <= 16 else t.rows[-1:]
        for row in presas:
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
        if fonte:
            self.legenda(f"Fonte: {fonte}", acima=False)
        return t

    def caixa(self, titulo, paragrafos, largura_cm=16.0, cor="F2F2F2"):
        """Caixa de destaque (tabela de 1 célula) que não se parte entre páginas."""
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _largura_fixa(t)
        t.columns[0].width = Cm(largura_cm)
        c = t.rows[0].cells[0]
        c.width = Cm(largura_cm)
        _sombrear(c, cor)
        _linha_nao_quebra(t.rows[0])
        par = c.paragraphs[0]
        par.style = self.doc.styles["Tabela"]
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(titulo.upper())
        r.bold = True
        r.font.size = Pt(11)
        for i, texto in enumerate(paragrafos):
            q = c.add_paragraph(style=self.doc.styles["Tabela"])
            q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            q.paragraph_format.space_after = Pt(4)
            q.paragraph_format.line_spacing = 1.15
            if i < len(paragrafos) - 1:
                q.paragraph_format.keep_with_next = True
            rr = q.add_run(texto)
            rr.font.size = Pt(11)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    def imagem(self, caminho, largura_cm, titulo, fonte):
        self.n_grafico += 1
        self.legenda(f"Gráfico {self.n_grafico} — {titulo}")
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(str(caminho), width=Cm(largura_cm))
        self.legenda(f"Fonte: {fonte}", acima=False)

    # ---------------- seções de página
    def secao(self, paisagem):
        """Nova seção começando em nova página (sem quebra extra = sem página em branco)."""
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        _margens(s, paisagem)
        self.paisagem = paisagem
        return s

    @property
    def largura_util(self):
        return 24.7 if self.paisagem else 16.0

    def salvar(self, caminho):
        self.doc.save(str(caminho))
        return caminho


# =============================================================================
# Utilidades de texto
# =============================================================================
def data_extenso(dt):
    return f"{dt.day} de {MESES_EXTENSO[dt.month - 1]} de {dt.year}"


def _fmt(v):
    return "—" if v is None else str(v)


def _fmt_dif(v):
    return fmt_dif(v)


def _objeto_curto(codigo):
    """Descrição curta a partir do código do arquivo (ex.: 'CNES · CBO 515105 · Set/2025')."""
    m = re.match(r"A_egestor_acs_\d+_(\d{6})_(\d{6})", codigo)
    if m:
        return f"e-Gestor · série ACS {rotulo_comp(m.group(1))} a {rotulo_comp(m.group(2))}"
    m = re.match(r"B_cnes_cbo(\d{6})_\d+_(\d{6})", codigo)
    if m:
        return f"CNES · CBO {m.group(1)} · {rotulo_comp(m.group(2))}"
    m = re.match(r"C_cnes_nominal_cbo(\d{6})", codigo)
    if m:
        return f"CNES · relação nominal CBO {m.group(1)}"
    if codigo.startswith("B0_"):
        return "CNES · formulário (competências disponíveis)"
    return codigo


def _hash_2_linhas(h):
    return [h[:32], h[32:]] if h else ["—"]


# =============================================================================
# Montagem do caderno
# =============================================================================
def gerar_docx(dados, caminho, graficos):
    mun = dados["municipio"]
    ini, fim = dados["inicio"], dados["fim"]
    quadro = dados["quadro"]
    setembros = [q for q in quadro if q["comp"].endswith("09")]
    outros = [q for q in quadro if not q["comp"].endswith("09")]
    ev = dados["evidencias"]
    ent = f"Município de {mun['nome']}/{mun['uf']}"
    fonte_cnes = "CNESWeb/DATASUS — consulta Profissionais por CBO; respostas preservadas (seção 2)."
    fonte_eg = "e-Gestor APS — Cobertura de ACS; resposta preservada (seção 2)."

    cd = Caderno()

    # ---------------------------------------------------------------- Abertura
    for txt, tam in (("CADERNO DE PROVAS", 14),
                     ("QUANTITATIVO DE AGENTES COMUNITÁRIOS DE SAÚDE (ACS) E DE AGENTES DE COMBATE ÀS ENDEMIAS (ACE)", 12),
                     (f"MUNICÍPIO DE {mun['nome'].upper()}/{mun['uf']}", 12)):
        par = cd.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(4)
        r = par.add_run(txt)
        r.bold = True
        r.font.size = Pt(tam)
    par = cd.doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run(f"Fontes: SCNES (CNESWeb/DATASUS) e e-Gestor APS · Período: {rotulo_comp(ini)} a {rotulo_comp(fim)} · "
                f"Extração: {dados['data_extracao_curta']}").font.size = Pt(10)

    # ---------------------------------------------------------------- 1
    cd.titulo("1 Objeto e identificação do ente")
    cd.p(f"O presente caderno reúne, em forma documental e reproduzível, o quantitativo de Agentes Comunitários "
         f"de Saúde (ACS) e de Agentes de Combate às Endemias (ACE) registrado nas bases públicas do Ministério "
         f"da Saúde para o {ent}, no período de {rotulo_comp(ini)} a {rotulo_comp(fim)}.")
    cd.p("Os números foram extraídos exclusivamente do Sistema de Cadastro Nacional de Estabelecimentos de Saúde "
         "(SCNES), por meio do CNESWeb/DATASUS, e do e-Gestor Atenção Primária à Saúde. Não há número estimado: "
         "o que não pôde ser extraído está declarado como não apurado.")
    cd.tabela(["Campo", "Informação"], [
        ["Ente", ent],
        ["Código IBGE (7 dígitos)", mun["ibge7"] or "—"],
        ["Código IBGE utilizado pelo CNES (6 dígitos)", mun["ibge6"]],
        ["Código da UF / código da região", f"{mun['couf']} / {mun['coregiao']}"],
        ["Período de referência", f"{rotulo_comp(ini)} a {rotulo_comp(fim)} (competências {ini} a {fim})"],
        ["Competências-chave do CNES", ", ".join(q["rotulo"] for q in quadro) or "—"],
        ["Última competência publicada — CNES", rotulo_comp(dados["ultima_cnes"]) if dados["ultima_cnes"] else "não apurada"],
        ["Última competência publicada — e-Gestor", rotulo_comp(dados["ultima_egestor"]) if dados["ultima_egestor"] else "não apurada"],
        ["Data e hora da extração", dados["data_extracao"]],
    ], [6.5, 9.5], titulo="Identificação do ente e do levantamento")

    # ---------------------------------------------------------------- 2
    cd.titulo("2 Metodologia, fontes e cadeia de obtenção")
    cd.titulo("2.1 Fontes consultadas e critério de apuração", 2)
    cd.p(f"Fonte A — e-Gestor APS. Série mensal do quantitativo de ACS, obtida em requisição única à interface "
         f"pública que alimenta o painel {URL_EGESTOR_PAINEL}, abrangendo todo o período. Campos utilizados: "
         f"nuComp (competência), qtAcsAtivoAb (ACS ativos), qtAcsCoberturaAb (ACS considerados no cálculo de "
         f"cobertura), qtPopulacao e pcCoberturaAcsAb.", negrito_ate="Fonte A — e-Gestor APS.")
    cd.p(f"Fonte B — CNESWeb/DATASUS. Consulta \"Profissionais por CBO\" ({URL_CNES_CBO}), por competência, para "
         f"os CBOs 515105 (Agente Comunitário de Saúde), 322255 (Técnico em Agente Comunitário de Saúde) e "
         f"515140 (Agente de Combate às Endemias), e a relação nominal correspondente a cada CBO.",
         negrito_ate="Fonte B — CNESWeb/DATASUS.")
    cd.p(f"Critério de apuração. O quantitativo de ACS corresponde à soma dos CBOs 515105 e 322255. O resultado foi "
         f"confrontado, competência a competência, com o campo qtAcsAtivoAb do e-Gestor; diferenças superiores a "
         f"{LIMITE_DIVERGENCIA} unidades, em módulo, estão registradas na seção 5. O e-Gestor não publica o "
         f"quantitativo de ACE; para ACE, a fonte é exclusivamente o CNES.", negrito_ate="Critério de apuração.")
    cd.titulo("2.2 Endereços e parâmetros exatos", 2)
    cd.p("As consultas seguiram quatro modelos de requisição. Os parâmetros que variam (competência e CBO) "
         "estão indicados no nome de cada arquivo preservado (seção 2.3).")
    eg_url = next((r["url"] for r in ev if r["codigo"].startswith("A_egestor")), "")
    cd.tabela(["Consulta (arquivos)", "Método", "Endereço exato", "Parâmetros enviados"], [
        ["Série mensal de ACS — e-Gestor (A_…)", "GET", eg_url, "contidos no próprio endereço"],
        ["Competências disponíveis — CNES (B0_…)", "GET", URL_CNES_CBO, "—"],
        ["Quantitativo por CBO — CNES (B_…)", "POST", URL_CNES_CBO,
         f"cboEstado={mun['couf']}; municipios={mun['ibge6']}; rdPesquisa=cbo; btnPesquisa=Pesquisar; "
         f"cboCompetencia=AAAAMM e cboOcupacao=CBO, conforme o nome do arquivo"],
        ["Relação nominal — CNES (C_…)", "GET",
         URL_CNES_NOMINAL.format(cbo="<CBO>", couf=mun["couf"], ibge6=mun["ibge6"]),
         "Vcbo = CBO indicado no nome do arquivo"],
    ], [3.4, 1.3, 6.3, 5.0], tamanho=8, alinhar="ECEE", titulo="Modelos de requisição utilizados")
    cd.titulo("2.3 Cadeia de obtenção e integridade", 2)
    cd.p("Cada resposta foi preservada em arquivo, exatamente como recebida do servidor, com cálculo do resumo "
         "criptográfico SHA-256. O resumo permite verificar que o arquivo não foi alterado após a extração "
         "(roteiro na seção 7). Os arquivos acompanham este caderno na pasta \"evidencias\".")
    linhas_ev = []
    for r in ev:
        obj = _objeto_curto(r["codigo"])
        if r["status"] != "ok":
            obj += f" — FALHA: {r['erro'][:60]}"
        dh = r["data_hora"].replace(" (horário de Brasília)", "").split(" ")
        linhas_ev.append([r["codigo"], obj, dh, _hash_2_linhas(r["sha256"])])
    cd.tabela(["Arquivo preservado (.json/.html)", "Objeto", "Data e hora (Brasília)", "SHA-256"], linhas_ev,
              [4.6, 3.4, 2.0, 6.0], tamanho=7.5, alinhar="EECE", mono_cols=(3,),
              titulo="Cadeia de obtenção — respostas preservadas e respectivos resumos SHA-256",
              fonte="elaboração própria a partir das respostas brutas preservadas.")

    # ---------------------------------------------------------------- 3
    cd.titulo("3 Quadro-síntese — competência de setembro de cada exercício")
    cd.p("O quadro apresenta, para a competência de setembro de cada exercício do período, o quantitativo "
         "registrado no CNES por CBO, o total de ACS (515105 + 322255), o quantitativo de ACE e o número de ACS "
         "ativos informado pelo e-Gestor na mesma competência. O sinal \"—\" indica valor não apurado "
         "(consulta sem resposta) ou competência ainda não publicada pela fonte.")
    cab3 = ["Exercício", "Compet.", "ACS 515105", "Técnico ACS 322255", "ACS total", "ACE 515140",
            "e-Gestor ACS ativos", "Diferença CNES − e-Gestor"]
    lin3 = [[q["comp"][:4], q["rotulo"], _fmt(q["cbo515105"]), _fmt(q["cbo322255"]), _fmt(q["acs_total"]),
             _fmt(q["ace515140"]), _fmt(q["egestor_ativo"]), _fmt_dif(q["diferenca"])] for q in setembros]
    if not lin3:
        lin3 = [["—"] * 8]
    cd.tabela(cab3, lin3, [1.9, 1.8, 1.8, 2.1, 1.7, 1.7, 2.4, 2.6], tamanho=10, alinhar="CCCCCCCC",
              titulo=f"Quantitativo de ACS e ACE — setembro de cada exercício — {ent}",
              fonte=f"{fonte_cnes} {fonte_eg} Elaboração própria.")
    if outros:
        cd.p("Competências adicionais consultadas (última publicada em cada base), para que a conferência fique "
             "pareada no mesmo mês:").paragraph_format.keep_with_next = True
        lin3b = [[q["rotulo"], q["papel"], _fmt(q["cbo515105"]), _fmt(q["cbo322255"]), _fmt(q["acs_total"]),
                  _fmt(q["ace515140"]), _fmt(q["egestor_ativo"]), _fmt_dif(q["diferenca"])] for q in outros]
        cd.tabela(["Compet.", "Papel", "ACS 515105", "Técnico ACS 322255", "ACS total", "ACE 515140",
                   "e-Gestor ACS ativos", "Diferença"], lin3b, [1.8, 2.7, 1.8, 2.1, 1.7, 1.7, 2.2, 2.0],
                  alinhar="CCCCCCCC", titulo="Última competência publicada em cada base",
                  fonte=f"{fonte_cnes} {fonte_eg}")

    # ---------------------------------------------------------------- 4 (paisagem)
    cd.secao(paisagem=True)
    cd.titulo("4 Série histórica mensal")
    pop = sorted({(l["anoPopulacao"], l["qtPopulacao"], l["origemPopulacao"]) for l in dados["egestor"]
                  if l["anoPopulacao"]}, key=lambda x: x[0])
    txt_pop = ""
    if pop:
        txt_pop = (" A população considerada pelo e-Gestor no cálculo de cobertura foi, por ano de referência: " +
                   "; ".join(f"{a}: {n:,} habitantes ({o})".replace(",", ".") for a, n, o in pop) + ".")
    cd.p("Quantitativo mensal de ACS ativos (campo qtAcsAtivoAb) informado pelo e-Gestor APS, organizado por "
         "exercício (linhas) e mês (colunas)." + txt_pop)
    mat = dados["matriz"]
    cd.tabela(["Exercício"] + MESES, [[str(a)] + [_fmt(v) for v in vals] for a, vals in mat],
              [2.3] + [1.86] * 12, tamanho=10, alinhar="C" * 13,
              titulo=f"ACS ativos por competência — e-Gestor APS — {ent}", fonte=fonte_eg)
    if graficos.get("serie"):
        cd.imagem(graficos["serie"], 17.0, "Série mensal do e-Gestor e totais do CNES nas competências-chave",
                  f"{fonte_eg} {fonte_cnes}")

    # ---------------------------------------------------------------- 5
    cd.secao(paisagem=False)
    cd.titulo("5 Composição por CBO")
    cd.p("O quadro abaixo decompõe o quantitativo de ACS entre os CBOs 515105 e 322255 em cada competência "
         "consultada e confronta o total com o e-Gestor. A coluna \"Situação\" registra \"Confere\" quando a "
         f"diferença é de até {LIMITE_DIVERGENCIA} unidades, em módulo.")
    lin5 = [[q["rotulo"], q["papel"] or "—", _fmt(q["cbo515105"]), _fmt(q["cbo322255"]), _fmt(q["acs_total"]),
             _fmt(q["ace515140"]), _fmt(q["egestor_ativo"]), _fmt_dif(q["diferenca"]), q["situacao"]]
            for q in quadro] or [["—"] * 9]
    cd.tabela(["Compet.", "Papel", "515105", "322255", "ACS total", "ACE 515140", "e-Gestor", "Dif.", "Situação"],
              lin5, [1.8, 2.5, 1.5, 1.5, 1.5, 1.5, 1.6, 1.4, 2.7], tamanho=9, alinhar="CCCCCCCCE",
              titulo="Composição do quantitativo de ACS por CBO e conferência com o e-Gestor",
              fonte=f"{fonte_cnes} {fonte_eg}")
    if graficos.get("composicao"):
        cd.imagem(graficos["composicao"], 16.0, "ACS por CBO nas competências-chave (515105 na base; 322255 no topo)",
                  fonte_cnes)
    ach = dados["achado"]
    cd.caixa("Achado relevante — reclassificação de CBO" if ach["ha_achado"] else
             "Verificação de reclassificação de CBO", ach["texto"])
    divs = [q for q in quadro if q["situacao"].startswith("Divergência")]
    if divs:
        cd.caixa("Divergências registradas para investigação",
                 [f"{q['rotulo']}: CNES (515105 + 322255) = {q['acs_total']}; e-Gestor (qtAcsAtivoAb) = "
                  f"{q['egestor_ativo']}; diferença = {_fmt_dif(q['diferenca'])}." for q in divs] +
                 ["Pontos de verificação, sem conclusão nesta peça: (i) datas de envio e processamento de cada base; "
                  "(ii) profissionais vinculados a estabelecimento fora de equipe da Atenção Primária; "
                  "(iii) retificação posterior do cadastro pelo gestor."], cor="FFF4E5")
    cd.p("Para ACE (CBO 515140), o e-Gestor não publica quantitativo; os números da coluna \"ACE 515140\" têm "
         "como fonte única o CNES.", tamanho=10, recuo=False)

    # ---------------------------------------------------------------- 6
    cd.titulo("6 Relação nominal por CBO")
    cd.p(f"Relações nominais extraídas do CNESWeb em {dados['data_extracao_curta']}. O CNESWeb não filtra a "
         f"relação nominal por competência pretérita: cada relação corresponde à posição cadastral da data da "
         f"consulta, e não às competências de setembro do quadro-síntese.")
    ult = {q["comp"]: q for q in quadro}.get(dados["ultima_cnes"]) if dados["ultima_cnes"] else None
    for k, (cbo, nom) in enumerate(dados["nominal"].items(), start=1):
        cd.titulo(f"6.{k} CBO {cbo} — {CBOS[cbo]}", 2)
        nomes = nom.get("nomes")
        if nomes is None:
            cd.p(f"Relação nominal não apurada ({nom.get('aviso', 'sem resposta')}).", recuo=False)
            continue
        ref_total = ult.get({"515105": "cbo515105", "322255": "cbo322255", "515140": "ace515140"}[cbo]) if ult else None
        frase = f"Total de nomes na relação: {len(nomes)}."
        if ref_total is not None:
            frase += f" Total informado pelo CNES na competência {ult['rotulo']} (última publicada): {ref_total}."
        cd.p(frase, recuo=False, tamanho=10)
        if not nomes:
            continue
        meio = math.ceil(len(nomes) / 2)
        esq, dir_ = nomes[:meio], nomes[meio:]
        linhas = []
        for i in range(meio):
            a = [str(i + 1), esq[i]]
            b = [str(meio + i + 1), dir_[i]] if i < len(dir_) else ["", ""]
            linhas.append(a + b)
        cd.tabela(["Nº", "Nome", "Nº", "Nome"], linhas, [1.0, 7.0, 1.0, 7.0], tamanho=9, alinhar="CECE",
                  titulo=f"Relação nominal — CBO {cbo} — posição de {dados['data_extracao_curta']}",
                  fonte=f"CNESWeb/DATASUS — relação nominal (arquivo {nom.get('evidencia', '')}).")

    # ---------------------------------------------------------------- 7
    cd.titulo("7 Roteiro de conferência")
    cd.p("Os passos abaixo permitem ao Juízo, às partes ou a perito reproduzir cada número deste caderno "
         "diretamente nas fontes públicas, sem necessidade de cadastro ou senha.")
    cd.titulo("7.1 Quantitativo por CBO (CNESWeb)", 2)
    cd.passos([
        f"Acessar {URL_CNES_CBO}.",
        f"Em \"Estado\", selecionar o código {mun['couf']} ({mun['uf']}); em \"Município\", selecionar "
        f"{mun['nome'].upper()} (código {mun['ibge6']}).",
        "Em \"Competência\", selecionar o mês desejado. A tela exibe a competência no formato MMAAAA "
        "(ex.: 092025 = setembro de 2025).",
        "Marcar a opção de pesquisa por CBO e selecionar a ocupação (515105, 322255 ou 515140); clicar em \"Pesquisar\".",
        "Conferir o número da coluna \"Total\" com a tabela da seção 5. Clicar no total abre a relação nominal "
        "(posição da data em que a consulta for feita).",
    ])
    cd.titulo("7.2 Série mensal de ACS (e-Gestor APS)", 2)
    cd.passos([
        f"Acessar {URL_EGESTOR_PAINEL}, selecionar unidade geográfica \"Município\", UF {mun['uf']} e "
        f"município {mun['nome']}, e a competência desejada (a tela exibe até 12 competências por vez).",
        "Conferir o campo de ACS ativos com a tabela da seção 4.",
        f"Alternativamente, abrir no navegador o endereço exato utilizado neste caderno, que devolve a série "
        f"completa: {eg_url}",
    ])
    cd.titulo("7.3 Integridade dos arquivos preservados", 2)
    cd.passos([
        "Abrir a pasta \"evidencias\" que acompanha este caderno.",
        "No Windows, executar no Prompt de Comando: certutil -hashfile \"NOME_DO_ARQUIVO\" SHA256. "
        "No Linux ou macOS: sha256sum NOME_DO_ARQUIVO.",
        "Comparar o resultado com a coluna SHA-256 da seção 2.3 (também listada em MANIFESTO_SHA256.txt). "
        "Resultados idênticos indicam que o arquivo não foi alterado após a extração.",
    ])

    # ---------------------------------------------------------------- 8
    cd.titulo("8 Nota metodológica, limites e relação de fontes")
    cd.titulo("8.1 Limites declarados", 2)
    cd.marcadores([
        "O CNES reflete o cadastro realizado pelo próprio gestor municipal; os números correspondem ao que foi "
        "informado à base, na data de cada competência.",
        "A relação nominal corresponde à posição cadastral da data da consulta; o CNESWeb não filtra a relação "
        "nominal por competência pretérita.",
        "O resumo SHA-256 atesta a integridade dos arquivos preservados. Nova consulta às fontes pode gerar arquivo "
        "com conteúdo binário distinto (por exemplo, atualização da página ou retificação do cadastro), sem prejuízo "
        "da conferência dos números de cada competência.",
        "O e-Gestor APS não publica quantitativo de ACE; para ACE, a fonte é exclusivamente o CNES.",
        "Os campos qtAcsAtivoAb (ACS ativos) e qtAcsCoberturaAb (ACS considerados no cálculo de cobertura) do "
        "e-Gestor podem diferir; a conferência com o CNES utiliza qtAcsAtivoAb.",
        "Os extratos de repasse financeiro (consultafns.saude.gov.br) não integram este caderno e serão juntados "
        "em apartado.",
    ])
    cd.titulo("8.2 Registros da coleta e pendências", 2)
    regs = dados["avisos"] + dados["pendencias"]
    if regs:
        cd.marcadores(regs)
    else:
        cd.p("Nenhuma ocorrência registrada durante a coleta: todas as consultas previstas retornaram resposta válida.")
    cd.titulo("8.3 Relação de fontes", 2)
    cd.marcadores([
        f"Ministério da Saúde. e-Gestor Atenção Primária à Saúde — Relatórios públicos: cobertura de ACS. "
        f"Disponível em: {URL_EGESTOR_PAINEL}. Acesso em: {dados['data_extracao_curta']}.",
        f"Ministério da Saúde. DATASUS. CNESWeb — Profissionais por CBO. Disponível em: {URL_CNES_CBO}. "
        f"Acesso em: {dados['data_extracao_curta']}.",
        f"Instituto Brasileiro de Geografia e Estatística. API de Localidades (códigos de municípios). "
        f"Disponível em: https://servicodados.ibge.gov.br/api/docs/localidades. Acesso em: {dados['data_extracao_curta']}.",
    ])

    cd.salvar(caminho)
    return caminho
